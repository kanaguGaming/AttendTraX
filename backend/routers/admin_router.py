# =============================================================================
# routers/admin_router.py  –  Admin-only endpoints
# =============================================================================
import io
import uuid
from typing import List, Optional

from fastapi import APIRouter, Depends, Form, HTTPException, UploadFile, File, status
from pydantic import BaseModel

import sheets
from auth import hash_password
from dependencies import require_role

router = APIRouter(prefix="/admin", tags=["Admin"])

_admin_dep = Depends(require_role("ADMIN"))


# ─────────────────────────────────────────────────────────────────────────────
# CLASSES
# ─────────────────────────────────────────────────────────────────────────────
class AddClassRequest(BaseModel):
    class_id: str
    class_name: str
    year: Optional[str] = ""
    section: Optional[str] = ""
    semester: Optional[str] = ""


@router.get("/classes")
async def list_all_classes(_=_admin_dep):
    return sheets.get_all_classes()


@router.post("/classes")
async def add_class(body: AddClassRequest, _=_admin_dep):
    sheets.add_class(body.model_dump())
    return {"message": f"Class '{body.class_name}' added."}


class RenameClassRequest(BaseModel):
    new_name: str


@router.patch("/classes/{class_id}/rename")
async def rename_class(class_id: str, body: RenameClassRequest, _=_admin_dep):
    sheets.update_class_name(class_id, body.new_name)
    return {"message": "Class renamed."}


@router.delete("/classes/{class_id}")
async def deactivate_class(class_id: str, _=_admin_dep):
    sheets.deactivate_class(class_id)
    return {"message": "Class deactivated."}


# ─────────────────────────────────────────────────────────────────────────────
# SUBJECTS
# ─────────────────────────────────────────────────────────────────────────────
class AddSubjectRequest(BaseModel):
    subject_id: str
    class_id: str
    subject_name: str
    faculty_id: Optional[str] = ""


@router.get("/classes/{class_id}/subjects")
async def list_subjects(class_id: str, _=_admin_dep):
    return sheets.get_subjects_for_class(class_id)


@router.post("/subjects")
async def add_subject(body: AddSubjectRequest, _=_admin_dep):
    sheets.add_subject(body.model_dump())
    return {"message": f"Subject '{body.subject_name}' added."}


@router.delete("/subjects/{subject_id}")
async def delete_subject(subject_id: str, _=_admin_dep):
    sheets.delete_subject(subject_id)
    return {"message": "Subject deactivated."}


# ─────────────────────────────────────────────────────────────────────────────
# STUDENTS
# ─────────────────────────────────────────────────────────────────────────────
class AddStudentRequest(BaseModel):
    reg_no: str
    name: str
    class_id: str
    class_name: Optional[str] = ""
    # Student login credentials
    username: Optional[str] = None   # defaults to reg_no
    password: Optional[str] = None   # defaults to reg_no


@router.get("/classes/{class_id}/students")
async def list_students(class_id: str, _=_admin_dep):
    return sheets.get_students_by_class(class_id)


@router.post("/students")
async def add_student(body: AddStudentRequest, _=_admin_dep):
    cls = sheets.get_class_by_id(body.class_id)
    class_name = body.class_name or (cls["ClassName"] if cls else body.class_id)

    # Add to Students master sheet
    sheets.add_student({
        "reg_no": body.reg_no,
        "name": body.name,
        "class_id": body.class_id,
        "class_name": class_name,
    })

    # Create login account
    username = body.username or body.reg_no
    password = body.password or body.reg_no
    sheets.create_user({
        "user_id": body.reg_no,
        "name": body.name,
        "username": username,
        "password_hash": hash_password(password),
        "role": "STUDENT",
        "class_id": body.class_id,
    })
    return {"message": f"Student '{body.name}' added with login '{username}'."}


@router.delete("/students/{reg_no}")
async def deactivate_student(reg_no: str, _=_admin_dep):
    sheets.deactivate_student(reg_no)
    return {"message": "Student deactivated."}


@router.post("/cache/refresh")
async def refresh_cache(_=_admin_dep):
    """
    Force-clears the in-memory read cache so the next request fetches
    fresh data from Google Sheets. Useful after manually editing the sheet.
    """
    sheets._cache_invalidate(
        "all_classes_raw",
        "all_students_raw",
        "all_subjects_raw",
        "all_users",
    )
    return {"message": "Cache cleared. Next read will fetch fresh data from Google Sheets."}



# ─────────────────────────────────────────────────────────────────────────────
# FACULTY ACCOUNTS
# ─────────────────────────────────────────────────────────────────────────────
class AddFacultyRequest(BaseModel):
    faculty_id: str
    name: str
    username: str
    password: str


@router.get("/faculty")
async def list_faculty(_=_admin_dep):
    users = sheets.get_all_users()
    # Filter for FACULTY role and return safe fields
    return [
        {
            "FacultyID": u.get("UserID", ""),
            "Name": u.get("Name", ""),
            "Username": u.get("Username", ""),
            "Role": u.get("Role", "")
        }
        for u in users if u.get("Role", "").upper() == "FACULTY"
    ]


@router.post("/faculty")
async def add_faculty(body: AddFacultyRequest, _=_admin_dep):
    sheets.create_user({
        "user_id": body.faculty_id,
        "name": body.name,
        "username": body.username,
        "password_hash": hash_password(body.password),
        "role": "FACULTY",
        "class_id": "",
    })
    return {"message": f"Faculty '{body.name}' added."}


@router.delete("/users/{username}")
async def delete_user(username: str, _=_admin_dep):
    sheets.delete_user(username)
    return {"message": f"User '{username}' deleted."}


# ─────────────────────────────────────────────────────────────────────────────
# REPORTS
# ─────────────────────────────────────────────────────────────────────────────
@router.get("/reports/class/{class_id}")
async def class_report(class_id: str, _=_admin_dep):
    return sheets.get_class_attendance_report(class_id)


@router.get("/reports/student/{reg_no}")
async def student_report(reg_no: str, _=_admin_dep):
    return sheets.get_student_attendance_summary(reg_no)


# ─────────────────────────────────────────────────────────────────────────────
# PAST ATTENDANCE IMPORT (Excel upload)
# ─────────────────────────────────────────────────────────────────────────────
@router.post("/import/attendance")
async def import_attendance(file: UploadFile = File(...), _=Depends(require_role("ADMIN"))):
    """
    Upload an .xlsx file with columns:
    Date(DD-MM-YYYY) | ClassID | Hour | SubjectID | RegNo | Status | FacultyID(optional)

    All rows are validated then bulk-inserted into Attendance_Log (skipping duplicates).
    Also writes to individual class sheets.
    """
    import openpyxl

    if not file.filename.endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="Only .xlsx files are supported.")

    contents = await file.read()
    wb = openpyxl.load_workbook(io.BytesIO(contents), data_only=True)
    ws_xl = wb.active

    headers = [str(c.value).strip() if c.value else "" for c in next(ws_xl.iter_rows(max_row=1))]
    required = {"Date", "ClassID", "Hour", "SubjectID", "RegNo", "Status"}
    missing = required - set(headers)
    if missing:
        raise HTTPException(
            status_code=400,
            detail=f"Missing required columns: {', '.join(missing)}. "
                   f"Expected headers: Date, ClassID, Hour, SubjectID, RegNo, Status, FacultyID",
        )

    idx = {h: i for i, h in enumerate(headers)}
    rows = []
    errors = []
    valid_statuses = {"P", "A", "-"}

    for row_num, row in enumerate(ws_xl.iter_rows(min_row=2, values_only=True), start=2):
        row_data = {h: (str(row[i]).strip() if row[i] is not None else "") for h, i in idx.items()}
        if not any(row_data.values()):
            continue   # skip empty rows

        # Basic validation
        if row_data.get("Status", "").upper() not in valid_statuses:
            errors.append(f"Row {row_num}: Invalid status '{row_data.get('Status')}'.")
            continue
        if row_data.get("Hour", "").upper() not in sheets.ALL_HOURS:
            errors.append(f"Row {row_num}: Invalid hour '{row_data.get('Hour')}'.")
            continue

        rows.append({
            "date": row_data["Date"],
            "class_id": row_data["ClassID"],
            "hour": row_data["Hour"].upper(),
            "subject_id": row_data["SubjectID"],
            "reg_no": row_data["RegNo"],
            "status": row_data["Status"].upper(),
            "faculty_id": row_data.get("FacultyID", "IMPORT"),
        })

    if errors:
        raise HTTPException(
            status_code=422,
            detail={"validation_errors": errors[:20], "message": "Fix errors and re-upload."},
        )

    inserted = sheets.import_past_attendance_rows(rows)
    return {
        "message": f"Import complete. {inserted} rows inserted (duplicates skipped).",
        "total_rows": len(rows),
        "inserted": inserted,
        "skipped": len(rows) - inserted,
    }


# ─────────────────────────────────────────────────────────────────────────────
# BULK IMPORT – STUDENTS  (Excel upload)
# ─────────────────────────────────────────────────────────────────────────────
@router.post("/import/students")
async def import_students(
    file: UploadFile = File(...),
    class_id: str = Form(...),           # selected from UI dropdown, NOT in Excel
    _=Depends(require_role("ADMIN")),
):
    """
    Upload an .xlsx file with two columns:
        col 1 – register number  ("register no" / "reg no" / "regno" / "roll no")
        col 2 – student name     ("student name" / "name")

    class_id  – passed as a form field (chosen from dropdown in the UI).
    username  – automatically set to the student's full name.
    password  – automatically set to the register number.

    Duplicate register numbers are silently skipped.
    """
    import openpyxl

    if not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="Only .xlsx files are supported.")

    if not class_id.strip():
        raise HTTPException(status_code=400, detail="Please select a class before uploading.")

    class_id   = class_id.strip()
    cls        = sheets.get_class_by_id(class_id)
    class_name = cls["ClassName"] if cls else class_id

    contents = await file.read()
    wb   = openpyxl.load_workbook(io.BytesIO(contents), data_only=True)
    ws_xl = wb.active

    raw_headers = [str(c.value).strip() if c.value else "" for c in next(ws_xl.iter_rows(max_row=1))]

    def _find_col(candidates, raw):
        """Case-insensitive, space/underscore-agnostic column finder."""
        norm = lambda s: s.lower().replace(" ", "").replace("_", "").replace(".", "")
        norm_raw = [norm(h) for h in raw]
        for c in candidates:
            if norm(c) in norm_raw:
                return norm_raw.index(norm(c))
        return -1

    reg_col  = _find_col(
        ["registerno", "regno", "reg", "rollno", "rollnumber", "registernum", "regnum"],
        raw_headers,
    )
    name_col = _find_col(
        ["studentname", "name", "fullname", "studentfullname", "sname"],
        raw_headers,
    )

    if reg_col == -1:
        raise HTTPException(
            status_code=400,
            detail=(
                'Column "Register No" not found. '
                'Accepted header names: "register no", "reg no", "regno", "roll no".'
            ),
        )
    if name_col == -1:
        raise HTTPException(
            status_code=400,
            detail=(
                'Column "Student Name" not found. '
                'Accepted header names: "student name", "name", "full name".'
            ),
        )

    # Existing reg nos → for duplicate detection
    existing_regnos = {
        str(s.get("RegNo", "")).strip().lower()
        for s in sheets.get_all_students()
    }

    inserted, skipped, errors = 0, 0, []

    for row_num, row in enumerate(ws_xl.iter_rows(min_row=2, values_only=True), start=2):
        reg_no = str(row[reg_col]).strip()  if row[reg_col]  is not None else ""
        name   = str(row[name_col]).strip() if row[name_col] is not None else ""

        if not reg_no and not name:
            continue          # completely blank row

        if not reg_no:
            errors.append(f"Row {row_num}: Register number is empty.")
            continue
        if not name:
            errors.append(f"Row {row_num}: Student name is empty (reg: {reg_no}).")
            continue

        if reg_no.lower() in existing_regnos:
            skipped += 1
            continue

        try:
            sheets.add_student({
                "reg_no":     reg_no,
                "name":       name,
                "class_id":   class_id,
                "class_name": class_name,
            })
            sheets.create_user({
                "user_id":       reg_no,
                "name":          name,
                "username":      name,        # username = student name
                "password_hash": hash_password(reg_no),  # password = register no
                "role":          "STUDENT",
                "class_id":      class_id,
            })
            existing_regnos.add(reg_no.lower())
            inserted += 1
        except Exception as ex:
            errors.append(f"Row {row_num} ({reg_no}): {str(ex)}")

    if errors:
        raise HTTPException(
            status_code=422,
            detail={
                "message": f"{inserted} added, {skipped} skipped. Errors found:",
                "errors":  errors[:20],
            },
        )

    return {
        "message":  f"Import complete. {inserted} student(s) added to '{class_name}', {skipped} already existed.",
        "inserted": inserted,
        "skipped":  skipped,
        "class":    class_name,
    }


# ─────────────────────────────────────────────────────────────────────────────
# BULK IMPORT – FACULTY  (Excel upload)
# ─────────────────────────────────────────────────────────────────────────────
@router.post("/import/faculty")
async def import_faculty(file: UploadFile = File(...), _=Depends(require_role("ADMIN"))):
    """
    Upload an .xlsx file to bulk-add faculty accounts.

    Required columns:  FacultyID | Name | Username | Password
    """
    import openpyxl

    if not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="Only .xlsx files are supported.")

    contents = await file.read()
    wb = openpyxl.load_workbook(io.BytesIO(contents), data_only=True)
    ws_xl = wb.active

    headers = [str(c.value).strip() if c.value else "" for c in next(ws_xl.iter_rows(max_row=1))]
    required = {"FacultyID", "Name", "Username", "Password"}
    missing = required - set(headers)
    if missing:
        raise HTTPException(
            status_code=400,
            detail=f"Missing columns: {', '.join(missing)}. Required: FacultyID, Name, Username, Password",
        )

    idx = {h: i for i, h in enumerate(headers)}

    # Existing usernames to detect duplicates
    existing_users = sheets.get_all_users()
    existing_usernames = {str(u.get("Username", "")).strip().lower() for u in existing_users}

    inserted, skipped, errors = 0, 0, []
    for row_num, row in enumerate(ws_xl.iter_rows(min_row=2, values_only=True), start=2):
        rd = {h: (str(row[i]).strip() if row[i] is not None else "") for h, i in idx.items()}
        if not any(rd.values()):
            continue

        faculty_id = rd.get("FacultyID", "").strip()
        name       = rd.get("Name", "").strip()
        username   = rd.get("Username", "").strip()
        password   = rd.get("Password", "").strip()

        if not faculty_id or not name or not username or not password:
            errors.append(f"Row {row_num}: All four columns (FacultyID, Name, Username, Password) are required.")
            continue

        if username.lower() in existing_usernames:
            skipped += 1
            continue

        try:
            sheets.create_user({
                "user_id": faculty_id,
                "name": name,
                "username": username,
                "password_hash": hash_password(password),
                "role": "FACULTY",
                "class_id": "",
            })
            existing_usernames.add(username.lower())
            inserted += 1
        except Exception as ex:
            errors.append(f"Row {row_num} ({username}): {str(ex)}")

    if errors:
        raise HTTPException(
            status_code=422,
            detail={"message": f"{inserted} inserted, {skipped} skipped. Errors found:", "errors": errors[:20]},
        )

    return {
        "message": f"Faculty import complete. {inserted} added, {skipped} skipped (username already exists).",
        "inserted": inserted,
        "skipped": skipped,
    }


# ─────────────────────────────────────────────────────────────────────────────
# BULK IMPORT – SUBJECTS  (Excel upload)
# ─────────────────────────────────────────────────────────────────────────────
@router.post("/import/subjects")
async def import_subjects(file: UploadFile = File(...), _=Depends(require_role("ADMIN"))):
    """
    Upload an .xlsx file to bulk-add subjects.

    Required columns:  SubjectID | SubjectName | ClassID
    Optional columns:  FacultyID
    """
    import openpyxl

    if not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="Only .xlsx files are supported.")

    contents = await file.read()
    wb = openpyxl.load_workbook(io.BytesIO(contents), data_only=True)
    ws_xl = wb.active

    headers = [str(c.value).strip() if c.value else "" for c in next(ws_xl.iter_rows(max_row=1))]
    required = {"SubjectID", "SubjectName", "ClassID"}
    missing = required - set(headers)
    if missing:
        raise HTTPException(
            status_code=400,
            detail=f"Missing columns: {', '.join(missing)}. Required: SubjectID, SubjectName, ClassID. Optional: FacultyID",
        )

    idx = {h: i for i, h in enumerate(headers)}

    inserted, skipped, errors = 0, 0, []
    for row_num, row in enumerate(ws_xl.iter_rows(min_row=2, values_only=True), start=2):
        rd = {h: (str(row[i]).strip() if row[i] is not None else "") for h, i in idx.items()}
        if not any(rd.values()):
            continue

        subject_id   = rd.get("SubjectID", "").strip()
        subject_name = rd.get("SubjectName", "").strip()
        class_id     = rd.get("ClassID", "").strip()
        faculty_id   = rd.get("FacultyID", "").strip()

        if not subject_id or not subject_name or not class_id:
            errors.append(f"Row {row_num}: SubjectID, SubjectName, and ClassID are required.")
            continue

        try:
            sheets.add_subject({
                "subject_id": subject_id,
                "class_id": class_id,
                "subject_name": subject_name,
                "faculty_id": faculty_id,
            })
            inserted += 1
        except Exception as ex:
            errors.append(f"Row {row_num} ({subject_id}): {str(ex)}")

    if errors:
        raise HTTPException(
            status_code=422,
            detail={"message": f"{inserted} inserted. Errors found:", "errors": errors[:20]},
        )

    return {
        "message": f"Subjects import complete. {inserted} subjects added.",
        "inserted": inserted,
        "skipped": skipped,
    }
